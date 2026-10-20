"""Build a plain engineering-report template from a default or one-level profile overlay."""
import argparse
from copy import deepcopy
import json
from pathlib import Path
import re

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

from report_io import read_json

DEFAULT = Path(__file__).resolve().parents[1] / 'assets/profiles/engineering-default/profile.json'
ALIGN = {'left': WD_ALIGN_PARAGRAPH.LEFT, 'center': WD_ALIGN_PARAGRAPH.CENTER,
         'right': WD_ALIGN_PARAGRAPH.RIGHT, 'justify': WD_ALIGN_PARAGRAPH.JUSTIFY}


def merge(base, overlay, location='profile'):
    for key, value in overlay.items():
        if key not in base:
            raise ValueError(f'unknown field: {location}.{key}')
        expected = base[key]
        if isinstance(expected, dict):
            if not isinstance(value, dict):
                raise ValueError(f'{location}.{key}: expected object')
            merge(expected, value, f'{location}.{key}')
        else:
            if expected is not None and not (
                type(value) is type(expected) or
                type(expected) in (int, float) and type(value) in (int, float)
            ):
                raise ValueError(f'{location}.{key}: wrong value type')
            base[key] = value
    return base


def load_profile(path=None):
    base = read_json(DEFAULT)
    if path and Path(path).resolve() != DEFAULT.resolve():
        overlay = read_json(path)
        if not isinstance(overlay, dict) or overlay.get('extends') != 'engineering-default':
            raise ValueError('custom profile must extend engineering-default')
        if not overlay.get('id') or overlay['id'] == base['id'] or 'label' not in overlay:
            raise ValueError('custom profile requires its own id and label')
        overlay = deepcopy(overlay)
        overlay.pop('extends')
        base = merge(base, overlay)
    if type(base['schema_version']) is not int or base['schema_version'] != 1:
        raise ValueError('unsupported schema_version')
    if not re.fullmatch(r'[a-z0-9]+(?:-[a-z0-9]+)*', base['id']):
        raise ValueError('profile id must use lowercase letters, numbers and hyphens')
    page = base['page']
    if min(page['width_mm'], page['height_mm'], page['margin_mm']) <= 0 or 2 * page['margin_mm'] >= min(page['width_mm'], page['height_mm']):
        raise ValueError('invalid page size or margin')
    for name, style in base['styles'].items():
        if style['alignment'] not in ALIGN or style['font_role'] not in {'body', 'heading'}:
            raise ValueError(f'{name}: unsupported alignment or font_role')
        if style['size_pt'] <= 0 or style['line'] <= 0 or min(style['before_pt'], style['after_pt']) < 0:
            raise ValueError(f'{name}: invalid size or spacing')
    if len(base['numbering']['headings']) != 3 or any(not isinstance(x, str) or not x for x in base['numbering']['headings']):
        raise ValueError('three nonempty heading patterns required')
    if base['table']['borders'] not in {'three-line', 'grid'}:
        raise ValueError('unsupported table borders')
    fill = base['table']['header_fill']
    if fill is not None and (not isinstance(fill, str) or not re.fullmatch('[0-9A-Fa-f]{6}', fill)):
        raise ValueError('header_fill must be null or six hex digits')
    if base['table']['cell_margin_mm'] < 0:
        raise ValueError('negative cell margin')
    for value in base['content'].values():
        if not isinstance(value, bool) and not (isinstance(value, list) and all(isinstance(x, str) for x in value)):
            raise ValueError('content values must be boolean or lists of section names')
    return base


def element(tag, **attrs):
    node = OxmlElement('w:' + tag)
    for key, value in attrs.items():
        node.set(qn('w:' + key), str(value))
    return node


def heading_numbering(doc, profile):
    root = doc.part.numbering_part.element
    abstract_id = max([int(n.get(qn('w:abstractNumId'))) for n in root.findall(qn('w:abstractNum'))] + [-1]) + 1
    num_id = max([int(n.get(qn('w:numId'))) for n in root.findall(qn('w:num'))] + [0]) + 1
    abstract = element('abstractNum', abstractNumId=abstract_id)
    abstract.append(element('multiLevelType', val='multilevel'))
    for level, pattern in enumerate(profile['numbering']['headings']):
        node = element('lvl', ilvl=level)
        for child in [element('start', val=1), element('numFmt', val='decimal'),
                      element('pStyle', val=doc.styles[f'Heading {level+1}'].style_id),
                      element('suff', val='space'), element('lvlText', val=pattern), element('lvlJc', val='left')]:
            node.append(child)
        abstract.append(node)
        style = doc.styles[f'Heading {level+1}']
        ppr = style.element.get_or_add_pPr()
        numpr = ppr.get_or_add_numPr()
        numpr.append(element('ilvl', val=level))
        numpr.append(element('numId', val=num_id))
        outline = ppr.find(qn('w:outlineLvl'))
        if outline is None:
            ppr.append(element('outlineLvl', val=level))
        else:
            outline.set(qn('w:val'), str(level))
    first_num = root.find(qn('w:num'))
    if first_num is None:
        root.append(abstract)
    else:
        first_num.addprevious(abstract)
    num = element('num', numId=num_id)
    num.append(element('abstractNumId', val=abstract_id))
    root.append(num)


def add_sequence(paragraph, kind):
    run = paragraph.add_run()
    run._r.append(element('fldChar', fldCharType='begin'))
    instruction = element('instrText')
    instruction.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
    instruction.text = f' SEQ {kind} \\* ARABIC '
    paragraph.add_run()._r.append(instruction)
    paragraph.add_run()._r.append(element('fldChar', fldCharType='separate'))
    paragraph.add_run('1')
    paragraph.add_run()._r.append(element('fldChar', fldCharType='end'))


def build(profile, output):
    page, content = profile['page'], profile['content']
    if page['header'] or page['page_numbers'] or any(content[k] for k in ['cover', 'toc', 'figure_list', 'table_list']):
        raise ValueError('formal page furniture requires retained-reference/host authoring; not supported by this builder')
    if profile['numbering']['captions'] != 'continuous' or profile['numbering']['equations'] != 'continuous':
        raise ValueError('chapter-based numbering requires host authoring')
    if profile['numbering']['items'] != ['（1）', '1）', '①']:
        raise ValueError('custom item ladder requires host authoring; profile may still be used for review')
    if output.exists():
        raise ValueError('refusing to overwrite an existing template')
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Mm(page['width_mm']), Mm(page['height_mm'])
    sec.top_margin = sec.bottom_margin = sec.left_margin = sec.right_margin = Mm(page['margin_mm'])
    sectpr = sec._sectPr
    for node in list(sectpr):
        if node.tag == qn('w:docGrid'):
            sectpr.remove(node)
    for name, spec in profile['styles'].items():
        style = doc.styles[name] if name in doc.styles else doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        if name != 'Normal':
            style.base_style = doc.styles['Normal']
        font = style.font
        font.name = profile['fonts']['latin']
        font.size, font.bold, font.italic, font.color.rgb = Pt(spec['size_pt']), spec['bold'], False, RGBColor(0, 0, 0)
        rf = style.element.get_or_add_rPr().get_or_add_rFonts()
        rf.set(qn('w:eastAsia'), profile['fonts']['heading' if spec['font_role'] == 'heading' else 'east_asia'])
        for key in ['asciiTheme', 'hAnsiTheme', 'eastAsiaTheme', 'cstheme']:
            rf.attrib.pop(qn('w:' + key), None)
        pf = style.paragraph_format
        pf.alignment = ALIGN[spec['alignment']]
        pf.line_spacing = spec['line']
        pf.space_before, pf.space_after = Pt(spec['before_pt']), Pt(spec['after_pt'])
        pf.first_line_indent, pf.left_indent = Pt(spec['first_line_pt']), Pt(spec['left_pt'])
        pf.keep_with_next, pf.widow_control = spec['keep_next'], True
        ppr = style.element.get_or_add_pPr()
        for inherited_decoration in list(ppr.findall(qn('w:pBdr'))):
            ppr.remove(inherited_decoration)
        ppr.append(element('snapToGrid', val=0))
        if name == 'Title':
            pf.keep_together = True
    heading_numbering(doc, profile)
    doc.add_paragraph(profile['label']+'模板', 'Title')
    doc.add_paragraph('使用说明', 'Report Lead')
    doc.add_paragraph('以下为格式示例，不是研究结果。正式报告应替换示例内容，删除使用说明；只保留适用章节，不为填满模板编造数据。')
    doc.add_paragraph('研究对象与目的', 'Heading 1')
    normal = profile['styles']['Normal']
    doc.add_paragraph(f'这里说明研究对象、问题、资料范围和目的。正文采用{profile["fonts"]["east_asia"]}与 {profile["fonts"]["latin"]}，{normal["size_pt"]:g} pt、{normal["line"]:g} 倍行距、首行缩进 {normal["first_line_pt"]:g} pt。')
    doc.add_paragraph('方法与评价指标', 'Heading 2')
    doc.add_paragraph('（1）分项标题。此处为解释性正文，整段不因分项编号而加粗。', 'Report Item')
    doc.add_paragraph('（2）后续分项。分项应结合整组内容检查编号和句末标点。', 'Report Item')
    doc.add_paragraph('适用条件', 'Heading 3')
    doc.add_paragraph('方法的条件、符号和单位在此解释。变量与单位分别设置字形，不对数学符号执行全局替换。')
    p = doc.add_paragraph('表 ', 'Report Table Caption')
    add_sequence(p, 'ReportTable')
    p.add_run(' 评价指标表格示例')
    table = doc.add_table(rows=3, cols=3)
    table.autofit = False
    width = (page['width_mm'] - 2*page['margin_mm']) / 3
    for col in table.columns:
        col.width = Mm(width)
    settings = profile['table']
    borders = element('tblBorders')
    for side in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        visible = settings['borders'] == 'grid' or side in {'top', 'bottom'}
        borders.append(element(side, val='single' if visible else 'nil', sz=8 if side in {'top', 'bottom'} else 4, color='000000'))
    table._tbl.tblPr.append(borders)
    margins = element('tblCellMar')
    for side in ['top', 'left', 'bottom', 'right']:
        margins.append(element(side, w=round(settings['cell_margin_mm']*1440/25.4), type='dxa'))
    table._tbl.tblPr.append(margins)
    values = [['指标', '定义或单位', '资料要求'], ['动作时间', 'ms', '真实记录'], ['相对误差', '%', '明确基准与分母']]
    for i, row in enumerate(table.rows):
        trpr = row._tr.get_or_add_trPr()
        if settings['split_rows'] is False:
            trpr.append(element('cantSplit'))
        if i == 0 and settings['repeat_header']:
            trpr.append(element('tblHeader'))
        for j, cell in enumerate(row.cells):
            cell.vertical_alignment = 1
            p = cell.paragraphs[0]
            p.style = doc.styles['Report Table Header' if i == 0 else ('Report Table Value' if j == 1 else 'Report Table Text')]
            p.add_run(values[i][j])
            if i == 0:
                tcpr = cell._tc.get_or_add_tcPr()
                line = element('tcBorders')
                line.append(element('bottom', val='single', sz=4, color='000000'))
                tcpr.append(line)
                if settings['header_fill']:
                    tcpr.append(element('shd', val='clear', fill=settings['header_fill']))
    doc.add_paragraph('注：以上仅演示表格格式，不代表实际测试数据。', 'Report Note')
    doc.add_paragraph('图像位置示例', 'Report Figure')
    p = doc.add_paragraph('图 ', 'Report Figure Caption')
    add_sequence(p, 'ReportFigure')
    p.add_run(' 图题在图下方')
    doc.add_paragraph('参考文献', 'Report Section')
    doc.add_paragraph('[1] 此处为参考文献格式槽位，正式成稿替换为核对过的真实条目。', 'Report Reference')
    doc.settings.element.append(element('updateFields', val='true'))
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--profile', type=Path)
    parser.add_argument('--output', type=Path, required=True)
    parser.add_argument('--resolved-out', type=Path)
    args = parser.parse_args()
    try:
        profile = load_profile(args.profile)
        if args.resolved_out and args.resolved_out.exists():
            raise ValueError('refusing to overwrite resolved profile')
        if args.resolved_out and args.resolved_out.resolve() == args.output.resolve():
            raise ValueError('output and resolved-out must differ')
        build(profile, args.output)
        if args.resolved_out:
            args.resolved_out.write_text(json.dumps(profile, ensure_ascii=False, indent=2), encoding='utf-8')
        print(json.dumps({'status': 'created', 'profile': profile['id'], 'output': str(args.output)}, ensure_ascii=False))
    except (OSError, ValueError, KeyError, TypeError) as exc:
        parser.exit(2, f'{exc}\n')


if __name__ == '__main__':
    main()
